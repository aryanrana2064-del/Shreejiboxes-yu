"""Aging / overdue exports and the statement business header."""
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from tests.helpers import ServiceTestCase
from vx7khata import exports, insights


def _pdf_text(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        return None
    return "\n".join(p.extract_text() for p in PdfReader(str(path)).pages)


class TestAgingAndOverdueExports(ServiceTestCase):
    def setUp(self):
        super().setUp()
        self.out = Path(self.tmp.name)
        a = self.svc.add_customer("Rahul", "9876543210")
        b = self.svc.add_customer("Amit")
        self.add(a.id, "UDHAAR", "01-11-2024", "Old chain", "200")
        self.add(a.id, "UDHAAR", "10-01-2025", "Ring", "1000", due_date="10-02-2025")
        self.add(a.id, "UDHAAR", "01-02-2025", "Coin", "500")
        self.add(a.id, "JAMA", "05-02-2025", "part", "300")
        self.add(b.id, "JAMA", "01-01-2025", "advance", "250")
        self.report = insights.aging_report(self.svc, as_of="20-02-2025")
        self.overdue = insights.overdue_items(self.svc, as_of="20-02-2025")

    def test_aging_xlsx(self):
        ws = load_workbook(exports.export_aging_xlsx(self.report, self.out / "a.xlsx", "My Shop"))["Aging"]
        rows = [[c.value for c in row] for row in ws.iter_rows()]
        rahul = next(r for r in rows if r[1] == "Rahul")
        self.assertEqual([round(float(v or 0) * 100) for v in rahul[3:9]], [50000, 90000, 0, 0, 140000, 0])
        self.assertEqual(rahul[9].date(), date(2025, 1, 10))
        amit = next(r for r in rows if r[1] == "Amit")
        self.assertEqual(round(float(amit[8]) * 100), 25000)  # advance
        total = next(r for r in rows if r[1] == "Total")
        self.assertEqual([round(float(v or 0) * 100) for v in total[3:9]], [50000, 90000, 0, 0, 140000, 25000])
        self.assertEqual(ws["B3"].value, "20-02-2025")

    def test_aging_pdf(self):
        path = exports.export_aging_pdf(self.report, self.out / "a.pdf", "My Shop")
        text = _pdf_text(path)
        if text is None:
            self.skipTest("pypdf not installed")
        for needle in ("Outstanding Aging Report", "20-02-2025", "Rahul", "Amit", "0-30 days", "90+ days",
                       "500.00", "900.00", "1,400.00", "250.00", "10-01-2025", "My Shop", "FIFO"):
            self.assertIn(needle, text, needle)

    def test_overdue_xlsx_and_pdf(self):
        self.assertEqual(len(self.overdue), 1)
        ws = load_workbook(exports.export_overdue_xlsx(self.overdue, "20-02-2025", self.out / "o.xlsx"))["Overdue"]
        rows = [[c.value for c in row] for row in ws.iter_rows()]
        row = next(r for r in rows if r[2] == "Ring")
        self.assertEqual((row[0], row[3].date(), row[4].date(), row[5]), ("Rahul", date(2025, 1, 10), date(2025, 2, 10), 10))
        self.assertEqual(round(float(row[6]) * 100), 90000)  # 1000 minus the 100 of payment that reached it
        text = _pdf_text(exports.export_overdue_pdf(self.overdue, "20-02-2025", self.out / "o.pdf"))
        if text is not None:
            for needle in ("Overdue Payments", "Ring", "10-02-2025", "900.00"):
                self.assertIn(needle, text, needle)

    def test_empty_reports_still_export(self):
        empty = insights.aging_report(self.svc, as_of="01-01-2020")
        self.assertTrue(exports.export_aging_pdf(empty, self.out / "e.pdf").exists())
        self.assertTrue(exports.export_aging_xlsx(empty, self.out / "e.xlsx").exists())
        self.assertTrue(exports.export_overdue_pdf([], "01-01-2020", self.out / "eo.pdf").exists())
        self.assertTrue(exports.export_overdue_xlsx([], "01-01-2020", self.out / "eo.xlsx").exists())


class TestStatementHeader(ServiceTestCase):
    def test_business_details_appear_on_statement(self):
        c = self.svc.add_customer("Rahul")
        self.add(c.id, "UDHAAR", "10-01-2025", "Chain", "100")
        led = self.svc.ledger(c.id)
        out = Path(self.tmp.name)
        ws = load_workbook(exports.export_statement_xlsx(
            led, out / "s.xlsx", "My Shop", "12 Main Road, Jaipur\nPh 98765 43210"))["Statement"]
        self.assertIn("12 Main Road, Jaipur Ph 98765 43210", ws["A2"].value)
        text = _pdf_text(exports.export_statement_pdf(led, out / "s.pdf", "My Shop", "12 Main Road, Jaipur\nPh 98765 43210"))
        if text is not None:
            self.assertIn("12 Main Road, Jaipur", text)
            self.assertIn("Ph 98765 43210", text)
        # old call style still works
        self.assertTrue(exports.export_statement_pdf(led, out / "s2.pdf", "My Shop").exists())
