"""PDF and Excel statements / reports.

Everything exported comes from the Ledger / Register / customer-summary objects
the user was looking at, so only the selected data is written.
"""
from __future__ import annotations

import os
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Optional, Sequence
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from . import _bundle_hints  # noqa: F401  (makes PyInstaller bundle ReportLab's dynamic imports)
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from . import APP_NAME, dates, money
from .labels import JAMA_LABEL, JAMA_WORD, UDHAAR_LABEL, UDHAAR_WORD
from .insights import BUCKET_LABELS, AgingReport, OverdueRow
from .service import CustomerSummary, Ledger, Register

# --------------------------------------------------------------------------- common
_FONTS: Optional[tuple[str, str, str]] = None


def _fonts() -> tuple[str, str, str]:
    """(regular, bold, currency symbol). Uses the bundled DejaVu Sans so the rupee sign prints."""
    global _FONTS
    if _FONTS is None:
        base = Path(__file__).resolve().parent / "assets" / "fonts"
        regular, bold = base / "DejaVuSans.ttf", base / "DejaVuSans-Bold.ttf"
        if regular.is_file() and bold.is_file():
            pdfmetrics.registerFont(TTFont("VX7Sans", str(regular)))
            pdfmetrics.registerFont(TTFont("VX7Sans-Bold", str(bold)))
            pdfmetrics.registerFontFamily("VX7Sans", normal="VX7Sans", bold="VX7Sans-Bold")
            _FONTS = ("VX7Sans", "VX7Sans-Bold", "₹")
        else:  # font files missing: still produce a correct PDF, with "Rs." instead of the rupee sign
            _FONTS = ("Helvetica", "Helvetica-Bold", "Rs.")
    return _FONTS


def _period_text(date_from: Optional[str], date_to: Optional[str]) -> str:
    if not date_from and not date_to:
        return "All dates"
    start = dates.format_date(date_from) if date_from else "Beginning"
    end = dates.format_date(date_to) if date_to else "Latest"
    return f"{start} to {end}"


def _qty_text(quantity: Optional[str]) -> str:
    return quantity or ""


def _rate_text(rate_paise: Optional[int]) -> str:
    return money.format_inr(rate_paise, symbol=False) if rate_paise else ""


def _stamp() -> str:
    return datetime.now().strftime("%d-%m-%Y %H:%M")


def _tmp_path(path: Path) -> Path:
    return path.with_name(path.name + ".part")


# --------------------------------------------------------------------------- PDF
def _pdf_doc(path: Path, title: str) -> SimpleDocTemplate:
    return SimpleDocTemplate(
        str(path), pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
        topMargin=12 * mm, bottomMargin=16 * mm, title=title, author=APP_NAME,
    )


def _pdf_footer(font: str):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(12 * mm, 8 * mm, f"Generated {_stamp()} by {APP_NAME}")
        canvas.drawRightString(landscape(A4)[0] - 12 * mm, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()
    return draw


def _styles(font: str, bold: str) -> dict[str, ParagraphStyle]:
    return {
        "title": ParagraphStyle("t", fontName=bold, fontSize=16, leading=20),
        "sub": ParagraphStyle("s", fontName=font, fontSize=9.5, leading=13, textColor=colors.HexColor("#333333")),
        "cell": ParagraphStyle("c", fontName=font, fontSize=8.5, leading=10.5),
        "cellb": ParagraphStyle("cb", fontName=bold, fontSize=8.5, leading=10.5),
    }


def _table_style(font: str, bold: str, numeric_cols: Sequence[int], emphasis_rows: Sequence[int]) -> TableStyle:
    cmds = [
        ("FONTNAME", (0, 0), (-1, -1), font),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f3a5f")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), bold),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#b8bfc9")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f6f9")]),
    ]
    for col in numeric_cols:
        cmds.append(("ALIGN", (col, 0), (col, -1), "RIGHT"))
    for row in emphasis_rows:
        cmds.append(("BACKGROUND", (0, row), (-1, row), colors.HexColor("#e3e9f2")))
        cmds.append(("FONTNAME", (0, row), (-1, row), bold))
    return TableStyle(cmds)


def _para(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(text or "").replace("\n", "<br/>"), style)


def _build_pdf(path: Path, doc: SimpleDocTemplate, story: list, font: str) -> Path:
    path = Path(path)
    tmp = _tmp_path(path)
    doc.filename = str(tmp)
    footer = _pdf_footer(font)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    os.replace(tmp, path)
    return path


def export_statement_pdf(ledger: Ledger, path, shop_name: str = "", shop_details: str = "") -> Path:
    font, bold, sym = _fonts()
    st = _styles(font, bold)
    path = Path(path)
    c = ledger.customer
    doc = _pdf_doc(path, f"Statement - {c.name}")

    story: list = [Paragraph(escape(shop_name or APP_NAME), st["title"])]
    if shop_details:
        story.append(Paragraph(escape(shop_details).replace("\n", "<br/>"), st["sub"]))
    story += [Paragraph("Customer Statement", st["sub"]), Spacer(1, 3 * mm)]
    details = [f"<b>Customer:</b> {escape(c.name)}"]
    if c.mobile:
        details.append(f"<b>Mobile:</b> {escape(c.mobile)}")
    if c.address:
        details.append(f"<b>Address:</b> {escape(c.address)}")
    details.append(f"<b>Period:</b> {_period_text(ledger.date_from, ledger.date_to)}")
    if ledger.search:
        details.append(f"<b>Item filter:</b> {escape(ledger.search)} "
                       "<i>(running balance is the true account balance)</i>")
    for line in details:
        story.append(Paragraph(line, st["sub"]))
    story.append(Spacer(1, 4 * mm))

    header = ["Date", "Item / Description", "Qty", f"Rate ({sym})", f"{UDHAAR_LABEL} ({sym})",
              f"{JAMA_LABEL} ({sym})", f"Balance ({sym})", "Notes"]
    data: list[list] = [header]
    data.append(["", _para("Opening balance", st["cellb"]), "", "", "", "",
                 money.format_inr(ledger.opening_balance, symbol=False), ""])
    for row in ledger.rows:
        t = row.txn
        data.append([
            dates.format_date(t.txn_date), _para(t.item, st["cell"]), _qty_text(t.quantity), _rate_text(t.rate_paise),
            money.format_inr(row.udhaar, symbol=False) if row.udhaar else "",
            money.format_inr(row.jama, symbol=False) if row.jama else "",
            money.format_inr(row.balance, symbol=False), _para(t.notes, st["cell"]),
        ])
    data.append(["", _para("Period totals / Closing balance", st["cellb"]), "", "",
                 money.format_inr(ledger.total_udhaar, symbol=False),
                 money.format_inr(ledger.total_jama, symbol=False),
                 money.format_inr(ledger.closing_balance, symbol=False), ""])

    widths = [24, 62, 18, 24, 28, 28, 31, 58]
    table = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
    table.setStyle(_table_style(font, bold, [2, 3, 4, 5, 6], [1, len(data) - 1]))
    story.append(table)
    story.append(Spacer(1, 4 * mm))

    summary = [
        ["Opening balance", money.format_inr(ledger.opening_balance)],
        ["Transactions in period", str(len(ledger.rows))],
        [f"Total {UDHAAR_WORD} (period)", money.format_inr(ledger.total_udhaar)],
        [f"Total {JAMA_WORD} (period)", money.format_inr(ledger.total_jama)],
        ["Closing balance", money.format_inr(ledger.closing_balance)],
    ]
    s_table = Table(summary, colWidths=[55 * mm, 40 * mm], hAlign="LEFT")
    s_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font), ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"), ("FONTNAME", (0, -1), (-1, -1), bold),
        ("LINEABOVE", (0, -1), (-1, -1), 0.8, colors.black),
        ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
    story.append(s_table)
    story.append(Spacer(1, 2 * mm))
    story.append(Paragraph("Positive balance = customer owes you; negative = advance / overpaid.", st["sub"]))
    return _build_pdf(path, doc, story, font)


def statement_html(ledger: Ledger, shop_name: str = "", shop_details: str = "") -> str:
    """Printable customer statement as simple HTML (rendered by Qt's text engine when printing)."""
    e = escape
    c = ledger.customer
    right = "align='right'"

    def cell(text, align="", bold=False):
        inner = f"<b>{e(text)}</b>" if bold else e(text)
        return f"<td {align} style='padding:3px;'>{inner}</td>"

    parts = [f"<body style='font-size:9pt;'><h2 style='margin:0;'>{e(shop_name or APP_NAME)}</h2>"]
    if shop_details:
        parts.append(f"<p style='margin:0;'>{e(shop_details).replace(chr(10), '<br/>')}</p>")
    parts.append("<p style='margin-top:6px;'><b>Customer Statement</b></p>")
    parts.append(f"<p style='margin:0;'><b>Customer:</b> {e(c.name)}</p>")
    if c.mobile:
        parts.append(f"<p style='margin:0;'><b>Mobile:</b> {e(c.mobile)}</p>")
    if c.address:
        parts.append(f"<p style='margin:0;'><b>Address:</b> {e(c.address)}</p>")
    parts.append(f"<p style='margin:0;'><b>Period:</b> {e(_period_text(ledger.date_from, ledger.date_to))}</p>")
    if ledger.search:
        parts.append(f"<p style='margin:0;'><b>Item filter:</b> {e(ledger.search)} "
                     "<i>(running balance is the true account balance)</i></p>")
    parts.append("<br/><table width='100%' border='1' cellspacing='0' cellpadding='2'>")
    heads = ["Date", "Item / Description", "Qty", "Rate", UDHAAR_LABEL, JAMA_LABEL, "Balance", "Notes"]
    parts.append("<thead><tr bgcolor='#dfe6f1'>" + "".join(f"<th style='padding:3px;'>{e(h)}</th>" for h in heads)
                 + "</tr></thead>")
    parts.append("<tr bgcolor='#eef2f8'>" + cell("") + cell("Opening balance", bold=True) + cell("") + cell("")
                 + cell("") + cell("") + cell(money.format_inr(ledger.opening_balance, symbol=False), right, True)
                 + cell("") + "</tr>")
    for row in ledger.rows:
        t = row.txn
        parts.append("<tr>" + cell(dates.format_date(t.txn_date)) + cell(t.item) + cell(_qty_text(t.quantity), right)
                     + cell(_rate_text(t.rate_paise), right)
                     + cell(money.format_inr(row.udhaar, symbol=False) if row.udhaar else "", right)
                     + cell(money.format_inr(row.jama, symbol=False) if row.jama else "", right)
                     + cell(money.format_inr(row.balance, symbol=False), right) + cell(t.notes or "") + "</tr>")
    parts.append("<tr bgcolor='#eef2f8'>" + cell("") + cell("Period totals / Closing balance", bold=True) + cell("")
                 + cell("") + cell(money.format_inr(ledger.total_udhaar, symbol=False), right, True)
                 + cell(money.format_inr(ledger.total_jama, symbol=False), right, True)
                 + cell(money.format_inr(ledger.closing_balance, symbol=False), right, True) + cell("") + "</tr>")
    parts.append("</table><br/>")
    parts.append("<table cellspacing='0' cellpadding='2'>")
    for label, value in (("Opening balance", money.format_inr(ledger.opening_balance)),
                         ("Transactions in period", str(len(ledger.rows))),
                         (f"Total {UDHAAR_WORD} (period)", money.format_inr(ledger.total_udhaar)),
                         (f"Total {JAMA_WORD} (period)", money.format_inr(ledger.total_jama)),
                         ("Closing balance", money.format_inr(ledger.closing_balance))):
        parts.append(f"<tr><td>{e(label)}</td><td align='right'>&nbsp;&nbsp;{e(value)}</td></tr>")
    parts.append("</table>")
    parts.append("<p><i>Positive balance = customer owes you; negative = advance / overpaid.</i></p>")
    parts.append(f"<p style='font-size:8pt;'>Printed on {_stamp()}</p></body>")
    return "".join(parts)


def export_register_pdf(register: Register, path, shop_name: str = "", customer_label: str = "All customers") -> Path:
    font, bold, sym = _fonts()
    st = _styles(font, bold)
    path = Path(path)
    doc = _pdf_doc(path, "Transaction Register")
    story: list = [Paragraph(escape(shop_name or APP_NAME), st["title"]),
                   Paragraph("Transaction Register", st["sub"]), Spacer(1, 3 * mm),
                   Paragraph(f"<b>Customers:</b> {escape(customer_label)}", st["sub"]),
                   Paragraph(f"<b>Period:</b> {_period_text(register.date_from, register.date_to)}", st["sub"])]
    if register.search:
        story.append(Paragraph(f"<b>Item filter:</b> {escape(register.search)}", st["sub"]))
    story.append(Spacer(1, 4 * mm))

    data: list[list] = [["Date", "Customer", "Item / Description", "Qty", f"Rate ({sym})",
                         f"{UDHAAR_LABEL} ({sym})", f"{JAMA_LABEL} ({sym})", "Notes"]]
    for row in register.rows:
        t = row.txn
        data.append([
            dates.format_date(t.txn_date), _para(row.customer_name, st["cell"]), _para(t.item, st["cell"]),
            _qty_text(t.quantity), _rate_text(t.rate_paise),
            money.format_inr(t.amount_paise, symbol=False) if t.txn_type == "UDHAAR" else "",
            money.format_inr(t.amount_paise, symbol=False) if t.txn_type == "JAMA" else "",
            _para(t.notes, st["cell"]),
        ])
    data.append(["", _para("Totals", st["cellb"]), "", "", "",
                 money.format_inr(register.total_udhaar, symbol=False),
                 money.format_inr(register.total_jama, symbol=False), ""])
    data.append(["", _para(f"Net ({UDHAAR_WORD} - {JAMA_WORD})", st["cellb"]), "", "", "",
                 money.format_inr(register.net, symbol=False), "", ""])
    widths = [24, 40, 56, 16, 24, 28, 28, 57]
    table = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
    table.setStyle(_table_style(font, bold, [3, 4, 5, 6], [len(data) - 2, len(data) - 1]))
    story.append(table)
    return _build_pdf(path, doc, story, font)


def export_customers_pdf(summaries: Sequence[CustomerSummary], path, shop_name: str = "", filter_text: str = "") -> Path:
    font, bold, sym = _fonts()
    st = _styles(font, bold)
    path = Path(path)
    doc = _pdf_doc(path, "Customer Balances")
    story: list = [Paragraph(escape(shop_name or APP_NAME), st["title"]),
                   Paragraph("Customer Balance Summary", st["sub"]), Spacer(1, 3 * mm),
                   Paragraph(f"<b>As on:</b> {datetime.now().strftime('%d-%m-%Y')}", st["sub"])]
    if filter_text:
        story.append(Paragraph(f"<b>Customer filter:</b> {escape(filter_text)}", st["sub"]))
    story.append(Spacer(1, 4 * mm))
    data: list[list] = [["#", "Customer", "Mobile", "Entries", f"Total {UDHAAR_LABEL} ({sym})",
                         f"Total {JAMA_LABEL} ({sym})", f"Balance ({sym})"]]
    tu = tj = 0
    for i, s in enumerate(summaries, 1):
        data.append([str(i), _para(s.customer.name, st["cell"]), s.customer.mobile, str(s.txn_count),
                     money.format_inr(s.total_udhaar, symbol=False), money.format_inr(s.total_jama, symbol=False),
                     money.format_inr(s.balance, symbol=False)])
        tu += s.total_udhaar
        tj += s.total_jama
    data.append(["", _para("Overall total", st["cellb"]), "", "", money.format_inr(tu, symbol=False),
                 money.format_inr(tj, symbol=False), money.format_inr(tu - tj, symbol=False)])
    widths = [10, 70, 36, 18, 45, 45, 45]
    table = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
    table.setStyle(_table_style(font, bold, [3, 4, 5, 6], [len(data) - 1]))
    story.append(table)
    return _build_pdf(path, doc, story, font)


# --------------------------------------------------------------------------- Excel
_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_EMPH_FILL = PatternFill("solid", fgColor="E3E9F2")
_THIN = Side(style="thin", color="B8BFC9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_MONEY_FMT = "#,##0.00"


def _text(ws, row: int, col: int, value: str, bold: bool = False, wrap: bool = False):
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        cell.data_type = "s"  # never let user text be interpreted as a formula
    if bold:
        cell.font = Font(bold=True)
    if wrap:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    return cell


def _money(ws, row: int, col: int, paise: Optional[int], bold: bool = False):
    cell = ws.cell(row=row, column=col)
    if paise is not None:
        cell.value = money.paise_to_decimal(paise)
        cell.number_format = _MONEY_FMT
    cell.alignment = Alignment(horizontal="right")
    if bold:
        cell.font = Font(bold=True)
    return cell


def _header_row(ws, row: int, headers: Sequence[str]) -> None:
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDER


def _style_range(ws, first_row: int, last_row: int, ncols: int, emphasis_rows: Sequence[int] = ()) -> None:
    for r in range(first_row, last_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = _BORDER
            if r in emphasis_rows:
                cell.fill = _EMPH_FILL
                cell.font = Font(bold=True)


def _widths(ws, widths: Sequence[float]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _save(wb: Workbook, path) -> Path:
    path = Path(path)
    tmp = _tmp_path(path)
    wb.save(str(tmp))
    os.replace(tmp, path)
    return path


def _qty_cell(ws, row: int, col: int, quantity: Optional[str]):
    cell = ws.cell(row=row, column=col)
    if quantity:
        cell.value = Decimal(quantity)
        cell.number_format = "0.###"
    cell.alignment = Alignment(horizontal="right")


def _date_cell(ws, row: int, col: int, iso: str):
    cell = ws.cell(row=row, column=col, value=dates.to_date(iso))
    cell.number_format = "DD-MM-YYYY"
    cell.alignment = Alignment(horizontal="center", vertical="top")


def export_statement_xlsx(ledger: Ledger, path, shop_name: str = "", shop_details: str = "") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    c = ledger.customer
    _text(ws, 1, 1, shop_name or APP_NAME, bold=True).font = Font(bold=True, size=14)
    _text(ws, 2, 1, "Customer Statement" + (f"  |  {' '.join(shop_details.split())}" if shop_details else ""))
    _text(ws, 3, 1, "Customer", bold=True)
    _text(ws, 3, 2, c.name)
    _text(ws, 4, 1, "Mobile", bold=True)
    _text(ws, 4, 2, c.mobile)
    _text(ws, 5, 1, "Address", bold=True)
    _text(ws, 5, 2, c.address)
    _text(ws, 6, 1, "Period", bold=True)
    _text(ws, 6, 2, _period_text(ledger.date_from, ledger.date_to))
    _text(ws, 7, 1, "Item filter", bold=True)
    _text(ws, 7, 2, ledger.search or "None")
    _text(ws, 8, 1, "Generated", bold=True)
    _text(ws, 8, 2, _stamp())

    head = 10
    _header_row(ws, head, ["Date", "Item / Description", "Quantity", "Rate (₹)", f"{UDHAAR_LABEL} (₹)",
                           f"{JAMA_LABEL} (₹)", "Balance (₹)", "Notes"])
    r = head + 1
    _text(ws, r, 2, "Opening balance", bold=True)
    _money(ws, r, 7, ledger.opening_balance, bold=True)
    opening_row = r
    for row in ledger.rows:
        r += 1
        t = row.txn
        _date_cell(ws, r, 1, t.txn_date)
        _text(ws, r, 2, t.item, wrap=True)
        _qty_cell(ws, r, 3, t.quantity)
        _money(ws, r, 4, t.rate_paise)
        _money(ws, r, 5, row.udhaar or None)
        _money(ws, r, 6, row.jama or None)
        _money(ws, r, 7, row.balance)
        _text(ws, r, 8, t.notes, wrap=True)
    r += 1
    _text(ws, r, 2, "Period totals / Closing balance", bold=True)
    _money(ws, r, 5, ledger.total_udhaar, bold=True)
    _money(ws, r, 6, ledger.total_jama, bold=True)
    _money(ws, r, 7, ledger.closing_balance, bold=True)
    _style_range(ws, head + 1, r, 8, emphasis_rows=[opening_row, r])

    r += 2
    for label, value in (("Opening balance", ledger.opening_balance), (f"Total {UDHAAR_WORD} (period)", ledger.total_udhaar),
                         (f"Total {JAMA_WORD} (period)", ledger.total_jama), ("Closing balance", ledger.closing_balance)):
        _text(ws, r, 2, label, bold=True)
        _money(ws, r, 3, value, bold=True)
        r += 1
    _text(ws, r, 2, "Positive balance = customer owes you; negative = advance / overpaid.")

    _widths(ws, [13, 38, 11, 13, 15, 15, 16, 36])
    ws.freeze_panes = ws.cell(row=head + 1, column=1)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{head}:{head}"
    return _save(wb, path)


def export_register_xlsx(register: Register, path, shop_name: str = "", customer_label: str = "All customers") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Register"
    _text(ws, 1, 1, shop_name or APP_NAME, bold=True).font = Font(bold=True, size=14)
    _text(ws, 2, 1, "Transaction Register")
    _text(ws, 3, 1, "Customers", bold=True)
    _text(ws, 3, 2, customer_label)
    _text(ws, 4, 1, "Period", bold=True)
    _text(ws, 4, 2, _period_text(register.date_from, register.date_to))
    _text(ws, 5, 1, "Item filter", bold=True)
    _text(ws, 5, 2, register.search or "None")
    _text(ws, 6, 1, "Generated", bold=True)
    _text(ws, 6, 2, _stamp())
    head = 8
    _header_row(ws, head, ["Date", "Customer", "Item / Description", "Quantity", "Rate (₹)",
                           f"{UDHAAR_LABEL} (₹)", f"{JAMA_LABEL} (₹)", "Notes"])
    r = head
    for row in register.rows:
        r += 1
        t = row.txn
        _date_cell(ws, r, 1, t.txn_date)
        _text(ws, r, 2, row.customer_name, wrap=True)
        _text(ws, r, 3, t.item, wrap=True)
        _qty_cell(ws, r, 4, t.quantity)
        _money(ws, r, 5, t.rate_paise)
        _money(ws, r, 6, t.amount_paise if t.txn_type == "UDHAAR" else None)
        _money(ws, r, 7, t.amount_paise if t.txn_type == "JAMA" else None)
        _text(ws, r, 8, t.notes, wrap=True)
    r += 1
    _text(ws, r, 2, "Totals", bold=True)
    _money(ws, r, 6, register.total_udhaar, bold=True)
    _money(ws, r, 7, register.total_jama, bold=True)
    r += 1
    _text(ws, r, 2, f"Net ({UDHAAR_WORD} - {JAMA_WORD})", bold=True)
    _money(ws, r, 6, register.net, bold=True)
    _style_range(ws, head + 1, r, 8, emphasis_rows=[r - 1, r])
    _widths(ws, [13, 24, 38, 11, 13, 15, 15, 36])
    ws.freeze_panes = ws.cell(row=head + 1, column=1)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return _save(wb, path)


def export_customers_xlsx(summaries: Sequence[CustomerSummary], path, shop_name: str = "", filter_text: str = "") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Customer balances"
    _text(ws, 1, 1, shop_name or APP_NAME, bold=True).font = Font(bold=True, size=14)
    _text(ws, 2, 1, "Customer Balance Summary")
    _text(ws, 3, 1, "As on", bold=True)
    _text(ws, 3, 2, datetime.now().strftime("%d-%m-%Y"))
    _text(ws, 4, 1, "Customer filter", bold=True)
    _text(ws, 4, 2, filter_text or "None")
    head = 6
    _header_row(ws, head, ["#", "Customer", "Mobile", "Entries", f"Total {UDHAAR_LABEL} (₹)",
                           f"Total {JAMA_LABEL} (₹)", "Balance (₹)"])
    r = head
    tu = tj = 0
    for i, s in enumerate(summaries, 1):
        r += 1
        ws.cell(row=r, column=1, value=i)
        _text(ws, r, 2, s.customer.name)
        _text(ws, r, 3, s.customer.mobile)
        ws.cell(row=r, column=4, value=s.txn_count)
        _money(ws, r, 5, s.total_udhaar)
        _money(ws, r, 6, s.total_jama)
        _money(ws, r, 7, s.balance)
        tu += s.total_udhaar
        tj += s.total_jama
    r += 1
    _text(ws, r, 2, "Overall total", bold=True)
    _money(ws, r, 5, tu, bold=True)
    _money(ws, r, 6, tj, bold=True)
    _money(ws, r, 7, tu - tj, bold=True)
    _style_range(ws, head + 1, r, 7, emphasis_rows=[r])
    _widths(ws, [6, 32, 18, 9, 18, 18, 18])
    ws.freeze_panes = ws.cell(row=head + 1, column=1)
    return _save(wb, path)


# --------------------------------------------------------------------------- aging & overdue
def _aging_data(report: AgingReport):
    header = ["#", "Customer", "Mobile", *BUCKET_LABELS, "Total due", "Advance", "Oldest unpaid"]
    rows = []
    for i, r in enumerate(report.rows, 1):
        rows.append([i, r.customer.name, r.customer.mobile, *r.buckets, r.total, r.advance, r.oldest_date])
    return header, rows


def export_aging_pdf(report: AgingReport, path, shop_name: str = "") -> Path:
    font, bold, sym = _fonts()
    st = _styles(font, bold)
    path = Path(path)
    doc = _pdf_doc(path, "Outstanding Aging")
    story: list = [Paragraph(escape(shop_name or APP_NAME), st["title"]),
                   Paragraph("Outstanding Aging Report", st["sub"]), Spacer(1, 3 * mm),
                   Paragraph(f"<b>As on:</b> {dates.format_date(report.as_of)}", st["sub"]),
                   Paragraph("Payments are applied to the oldest unpaid credit first (FIFO). "
                             "Age is counted from the credit's transaction date.", st["sub"]),
                   Spacer(1, 4 * mm)]
    header, rows = _aging_data(report)
    data: list[list] = [[h if i < 3 or i == 9 else f"{h} ({sym})" for i, h in enumerate(header)]]
    for r in rows:
        data.append([str(r[0]), _para(r[1], st["cell"]), r[2],
                     *[money.format_inr(v, symbol=False) if v else "" for v in r[3:8]],
                     money.format_inr(r[8], symbol=False) if r[8] else "",
                     dates.format_date(r[9]) if r[9] else ""])
    t = report.bucket_totals
    data.append(["", _para("Total", st["cellb"]), "", *[money.format_inr(v, symbol=False) for v in t],
                 money.format_inr(report.total_outstanding, symbol=False),
                 money.format_inr(report.total_advance, symbol=False), ""])
    widths = [8, 52, 28, 26, 26, 26, 26, 28, 22, 24]
    table = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
    table.setStyle(_table_style(font, bold, [3, 4, 5, 6, 7, 8], [len(data) - 1]))
    story.append(table)
    return _build_pdf(path, doc, story, font)


def export_aging_xlsx(report: AgingReport, path, shop_name: str = "") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Aging"
    _text(ws, 1, 1, shop_name or APP_NAME, bold=True).font = Font(bold=True, size=14)
    _text(ws, 2, 1, "Outstanding Aging Report")
    _text(ws, 3, 1, "As on", bold=True)
    _text(ws, 3, 2, dates.format_date(report.as_of))
    _text(ws, 4, 1, "Method", bold=True)
    _text(ws, 4, 2, "Payments applied to the oldest unpaid credit first (FIFO); age from transaction date")
    header, rows = _aging_data(report)
    head = 6
    _header_row(ws, head, [h if i < 3 or i == 9 else f"{h} (\u20b9)" for i, h in enumerate(header)])
    r = head
    for row in rows:
        r += 1
        ws.cell(row=r, column=1, value=row[0])
        _text(ws, r, 2, row[1])
        _text(ws, r, 3, row[2])
        for col, v in zip(range(4, 9), row[3:8]):
            _money(ws, r, col, v or None)
        _money(ws, r, 9, row[8] or None)
        if row[9]:
            _date_cell(ws, r, 10, row[9])
    r += 1
    _text(ws, r, 2, "Total", bold=True)
    for col, v in zip(range(4, 8), report.bucket_totals):
        _money(ws, r, col, v, bold=True)
    _money(ws, r, 8, report.total_outstanding, bold=True)
    _money(ws, r, 9, report.total_advance, bold=True)
    _style_range(ws, head + 1, r, 10, emphasis_rows=[r])
    _widths(ws, [6, 30, 16, 15, 15, 15, 15, 16, 14, 16])
    ws.freeze_panes = ws.cell(row=head + 1, column=1)
    return _save(wb, path)


def export_overdue_pdf(rows: Sequence[OverdueRow], as_of: str, path, shop_name: str = "") -> Path:
    font, bold, sym = _fonts()
    st = _styles(font, bold)
    path = Path(path)
    doc = _pdf_doc(path, "Overdue Payments")
    story: list = [Paragraph(escape(shop_name or APP_NAME), st["title"]),
                   Paragraph("Overdue Payments", st["sub"]), Spacer(1, 3 * mm),
                   Paragraph(f"<b>As on:</b> {dates.format_date(as_of)}", st["sub"]), Spacer(1, 4 * mm)]
    data: list[list] = [["Customer", "Mobile", "Item / Description", "Entry date", "Due date", "Days overdue",
                         f"Unpaid ({sym})"]]
    for r in rows:
        data.append([_para(r.customer.name, st["cell"]), r.customer.mobile, _para(r.txn.item, st["cell"]),
                     dates.format_date(r.txn.txn_date), dates.format_date(r.txn.due_date), str(r.days_overdue),
                     money.format_inr(r.remaining, symbol=False)])
    data.append([_para("Total", st["cellb"]), "", "", "", "", "", money.format_inr(sum(r.remaining for r in rows), symbol=False)])
    widths = [48, 30, 70, 24, 24, 22, 40]
    table = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
    table.setStyle(_table_style(font, bold, [5, 6], [len(data) - 1]))
    story.append(table)
    return _build_pdf(path, doc, story, font)


def export_overdue_xlsx(rows: Sequence[OverdueRow], as_of: str, path, shop_name: str = "") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Overdue"
    _text(ws, 1, 1, shop_name or APP_NAME, bold=True).font = Font(bold=True, size=14)
    _text(ws, 2, 1, "Overdue Payments")
    _text(ws, 3, 1, "As on", bold=True)
    _text(ws, 3, 2, dates.format_date(as_of))
    head = 5
    _header_row(ws, head, ["Customer", "Mobile", "Item / Description", "Entry date", "Due date", "Days overdue",
                           "Unpaid (\u20b9)"])
    r = head
    for row in rows:
        r += 1
        _text(ws, r, 1, row.customer.name)
        _text(ws, r, 2, row.customer.mobile)
        _text(ws, r, 3, row.txn.item, wrap=True)
        _date_cell(ws, r, 4, row.txn.txn_date)
        _date_cell(ws, r, 5, row.txn.due_date)
        ws.cell(row=r, column=6, value=row.days_overdue)
        _money(ws, r, 7, row.remaining)
    r += 1
    _text(ws, r, 1, "Total", bold=True)
    _money(ws, r, 7, sum(x.remaining for x in rows), bold=True)
    _style_range(ws, head + 1, r, 7, emphasis_rows=[r])
    _widths(ws, [28, 16, 36, 14, 14, 14, 16])
    ws.freeze_panes = ws.cell(row=head + 1, column=1)
    return _save(wb, path)
